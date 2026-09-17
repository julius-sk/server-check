"""Build the review spreadsheet from whatever the audit script printed.

Input:  outputs/<host>.txt  -- one file per server, either the console capture or
        the report_<host>_<ts>.txt the script writes. Both work.
Output: server_security_review_<n>servers_<ts>.xlsx, same four-sheet shape as the
        previous review: Summary, one sheet per host, Findings & Actions,
        Evidence Commands.

Two things about the input that the parser has to survive:

  * TERMINAL WRAPPING. A console capture breaks long lines at the terminal width
    and pads the remainder with spaces, so "restrict to the" and "lab subnet."
    arrive on different lines. Any line that does not begin with a known marker
    is therefore treated as a continuation of the one before it, and runs of
    whitespace are collapsed.
  * DUPLICATE FINDINGS. The same condition is reported once per listening socket,
    so NFS on 2049 and 111 appears four times on a dual-stack host. Findings are
    de-duplicated on their text before they reach the Findings sheet.

Risk and severity are assigned by matching the finding text against the table in
RULES. That table encodes judgements a human made during the previous review, so
it is the part to edit when the judgement changes -- not the per-host output.

Usage:
    python make_report.py [outputs_dir] [-o out.xlsx]
"""
import argparse
import datetime as dt
import glob
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MARKERS = ("[OK]", "[WARN]", "[FAIL]", "[info]", "===", "ALL IFACES", "IFACE",
           "Listening TCP", "UID >= 1000", "Failures by", "Findings (", "Host:",
           "No  Check", "Legend:", "Report:", "Summary:", "Evidence:", "Reminder:",
           "Server Security Review", "Date:")

ITEMS = [
    ("DMZ Server Usage Status Review", "- Verify active usage"),
    ("DMZ Server Usage Status Review", "- Check active services (ports) and communication status"),
    ("Server Account Status Review", "- Check for unauthorized accounts (review general user accounts)"),
    ("Server Account Status Review", "- Check for long-term unaccessed (idle) accounts"),
    ("Server Security Log Review", "- Check for abnormal access attempts"),
    ("Command History Review", "- Check for suspicious command execution history"),
]

# (regex on the finding text, risk, severity, recommended action)
RULES = [
    (r"never logged in but still enabled", "Dormant credentials can be reused or brute-forced without the owner noticing.",
     "High", "Lock each account (usermod -L; chage -E 0), confirm project status with the owner, and delete if the work has ended."),
    (r"idle for more than", "Dormant credentials can be reused or brute-forced without the owner noticing.",
     "High", "Lock each account (usermod -L; chage -E 0), confirm project status with the owner, and delete if the work has ended."),
    (r"has NO PASSWORD set", "An account with no password can be entered directly.",
     "High", "Set a password or lock the account immediately."),
    (r"non-root account\(s\) with UID 0", "A second UID 0 account is a hidden root backdoor.",
     "High", "Identify who created it, remove it, and audit what it was used for."),
    (r"legacy cleartext service", "Credentials and session data cross the network in the clear.",
     "High", "Disable the service; use SSH or an encrypted equivalent."),
    (r"PermitRootLogin=yes", "Direct root SSH login removes attribution and is a standing brute-force target.",
     "High", "Set PermitRootLogin=no and use named accounts with sudo."),
    (r"no_root_squash|wildcard host", "Any client that can reach the share obtains local root on it.",
     "High", "Restrict /etc/exports to the lab subnet and enable root_squash."),
    (r"ld\.so\.preload exists", "Classic rootkit hook.",
     "High", "Investigate immediately; treat the host as suspect until cleared."),
    (r"PasswordAuthentication=yes", "Password logins on an internet-reachable host are directly exploitable, especially with shared or weak passwords.",
     "Medium", "Move to key-only authentication (PasswordAuthentication no) with a second session open, then rotate the account passwords."),
    (r"NFS/rpcbind exposed", "Unauthenticated share enumeration or mount from outside the lab subnet.",
     "Medium", "Restrict /etc/exports to the lab subnet, verify root_squash, and firewall 111/2049."),
    (r"metrics/dashboard port", "Host inventory, hardware and workload metrics readable by anyone who can reach the port.",
     "Medium", "Bind the exporter or dashboard to localhost, or firewall the port to the management network."),
    (r"web/dev application port", "A development web server without TLS or strong authentication, reachable from the network.",
     "High", "Put it behind a reverse proxy with TLS and authentication, or bind it to localhost."),
    (r"proxy port", "A reachable proxy can be used to bypass network egress controls.",
     "Medium", "Confirm the owning process and bind it to localhost if it is a personal tunnel."),
    (r"has no password expiry set", "Passwords never age out, so a leaked credential stays valid indefinitely.",
     "Medium", "Set an expiry policy (chage -M 180 -W 14) on every account that stays."),
    (r"commands worth explaining", "Commands in this class can install persistence, suppress logging, or move data between projects. Each hit needs a human explanation.",
     "Medium", "Review every listed line in evidence/06_history.txt and record why it was run."),
    (r"bash_history is a SYMLINK", "History suppression.",
     "High", "Investigate who redirected it and what was run in that period."),
    (r"bash_history does not exist", "Item 6 cannot be evidenced from root history.",
     "Medium", "Check /root/ for deletion or a symlink, and use sudo entries in auth.log as the alternative evidence."),
    (r"process names unavailable without root", "Listeners cannot be attributed, so item 2 is incomplete.",
     "Medium", "Re-run the audit with sudo."),
    (r"failed login for non-existent user", "An unregistered account name or a shared workstation is attempting access.",
     "Low", "Identify the source host and its owner, then confirm whether the attempt was legitimate."),
    (r"failed logins - check for brute-force", "Sustained failed logins from one source indicate a brute-force attempt.",
     "High", "Identify the source, block it, and enable fail2ban."),
    (r"empty GECOS field", "Accountability for actions taken by the account cannot be established.",
     "Low", "Fill the GECOS field with the owner's real name and affiliation, and attach the approval record."),
    (r"btmp only covers", "Insufficient forensic evidence for the full audit period.",
     "Low", "Extend logrotate retention for btmp, wtmp and auth.log to cover 12 months."),
    (r"HISTTIMEFORMAT is not set", "Shell history has no timestamps, so actions cannot be placed in time.",
     "Low", "Set HISTTIMEFORMAT and a larger HISTSIZE in /etc/profile.d/."),
    (r"UID gap", "Undocumented account lifecycle; a home directory or SSH key may remain.",
     "Low", "Confirm the deletion was approved and that the home directory and keys were removed."),
    (r"files owned by a deleted user", "Leftover data from a removed account.",
     "Low", "Reassign or delete the files."),
]
SEV_ORDER = {"High": 0, "Medium": 1, "Low": 2}
SEV_DAYS = {"High": 7, "Medium": 21, "Low": 45}

# Repairing terminal damage.
#
# A console capture pads each wrapped line to the terminal width with a long run
# of spaces, and that run lands INSIDE the logical line rather than starting a
# new one. Collapsing it is lossy: "restrict to the<pad>lab subnet" needs a space
# restored, "PasswordAuthenticati<pad>on=yes" needs none, and nothing in the text
# says which.
#
# It does not have to be guessed, because the audit script emits a closed set of
# messages. Matching a finding against those templates with every space removed
# identifies it regardless of where the terminal broke it, and the canonical text
# is then re-rendered from the template. Findings read from an unwrapped
# report_*.txt pass through unchanged.
CANON = [
    (r"legacycleartextserviceonport(\d+)\((.+?)\)",
     "2: legacy cleartext service on port {0} ({1}) - must be disabled."),
    (r"NFS/rpcbindexposedonallinterfaces\(port(\d+)\)",
     "2: NFS/rpcbind exposed on all interfaces (port {0}) - restrict to the lab subnet."),
    (r"metrics/dashboardport(\d+)openwithoutauthentication",
     "2: metrics/dashboard port {0} open without authentication - bind to localhost or firewall it."),
    (r"web/devapplicationport(\d+)opentoallinterfaces",
     "2: web/dev application port {0} open to all interfaces - put behind a reverse proxy with TLS+auth or restrict it."),
    (r"proxyport(\d+)reachablefromthenetwork",
     "2: proxy port {0} reachable from the network - confirm it is not an egress bypass."),
    (r"TelnetorFTPisrunning",
     "2: Telnet or FTP is running."),
    (r"processnamesunavailablewithoutroot",
     "2: process names unavailable without root - re-run with sudo to attribute every listener."),
    (r"PermitRootLogin=yes",
     "2: PermitRootLogin=yes - direct root SSH login must be disabled."),
    (r"PasswordAuthentication=yes",
     "2: PasswordAuthentication=yes - shared/weak passwords are directly exploitable; move to key-only auth."),
    (r"/etc/exportscontainsawildcardhost",
     "2: /etc/exports contains a wildcard host (*) - any client can mount."),
    (r"/etc/exportsusesno_root_squash",
     "2: /etc/exports uses no_root_squash - remote root gets local root on the share."),
    (r"account'([^']+)'hasanemptyGECOSfield",
     "3: account '{0}' has an empty GECOS field - record the owner's real name and affiliation."),
    (r"account'([^']+)'hasNOPASSWORDset",
     "3: account '{0}' has NO PASSWORD set."),
    (r"non-rootaccount\(s\)withUID0:(.+)$",
     "3: non-root account(s) with UID 0: {0}"),
    (r"UIDgap\(s\)(.+?)-anaccountwasprobablydeleted",
     "3: UID gap(s) {0} - an account was probably deleted; confirm it was approved and that its home directory and SSH keys were removed."),
    (r"filesownedbyadeleteduserremainunder/home:(.+)$",
     "3: files owned by a deleted user remain under /home: {0}"),
    (r"account'([^']+)'hasnopasswordexpiryset",
     "4: account '{0}' has no password expiry set."),
    (r"account\(s\)neverloggedinbutstillenabled:(.+)$",
     "4: account(s) never logged in but still enabled: {0}"),
    (r"account\(s\)idleformorethan(\d+)days:(.+)$",
     "4: account(s) idle for more than {0} days: {1}"),
    (r"failedloginfornon-existentuser'([^']+)'",
     "5: failed login for non-existent user '{0}' - identify the source host/owner."),
    (r"(\d+)failedlogins-checkforbrute-force",
     "5: {0} failed logins - check for brute-force from a single source."),
    (r"(\d+)'Invaliduser'entriesinauth\.log",
     "5: {0} 'Invalid user' entries in auth.log - review the sources."),
    (r"btmponlycovers~(\d+)days",
     "5: btmp only covers ~{0} days - extend log retention to cover the audit period."),
    (r"bash_historyisaSYMLINK\((.+?)\)",
     "6: /root/.bash_history is a SYMLINK ({0}) - history suppression; investigate."),
    (r"bash_historydoesnotexist",
     "6: /root/.bash_history does not exist - item 6 cannot be evidenced from root history; use sudo entries in auth.log instead."),
    (r"commandsworthexplainingwerefoundinshellhistory",
     "6: commands worth explaining were found in shell history (see evidence/06_history.txt)."),
    (r"HISTTIMEFORMATisnotset",
     "6: HISTTIMEFORMAT is not set - shell history has no timestamps."),
    (r"(/etc/ld\.so\.preload)exists",
     "6: {0} exists - classic rootkit hook, investigate immediately."),
]


def canonical(text):
    """Re-render a finding from its template, undoing terminal wrap damage."""
    flat = re.sub(r"\s+", "", text)
    for rx, tpl in CANON:
        m = re.search(rx, flat, re.I)
        if m:
            groups = list(m.groups())
            # Templates ending in a free-text tail carry a space-separated list of
            # account names. Those spaces were destroyed by the flattening used to
            # identify the template, so the tail is taken from the original text
            # instead, where only runs of whitespace were collapsed.
            if rx.endswith(r"(.+)$") and groups:
                groups[-1] = re.sub(r"\s+", " ", text).rsplit(":", 1)[-1].strip()
            try:
                return tpl.format(*groups)
            except (IndexError, KeyError):
                return tpl
    return text


ANSI = re.compile(r"\x1b\[[0-9;]*[A-Za-z]")


def unwrap(text):
    """Rejoin terminal-wrapped lines and collapse padding whitespace.

    Colour codes are stripped first. The audit script writes the report with tee
    while stdout is a terminal, so the saved file carries the escape sequences
    too; left in place they break line-start detection (a line beginning with an
    escape does not start with "[WARN]") and openpyxl refuses to write them.
    """
    text = ANSI.sub("", text)
    out = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            out.append("")
            continue
        stripped = line.strip()
        starts = any(stripped.startswith(m) for m in MARKERS) or \
            re.match(r"^\d+\.\s", stripped) or re.match(r"^\d+\s{2,}", stripped) or \
            re.match(r"^\d+:", stripped) or re.match(r"^\d+\s+\S+$", stripped)
        if out and out[-1] and not starts:
            out[-1] = out[-1].rstrip() + " " + stripped
        else:
            out.append(stripped)
    return [re.sub(r"\s{2,}", " ", l).strip() for l in out]


def read_source(path):
    """Return report text, preferring a pasted REPORT PACK over console scrollback.

    The pack is base64 of a gzipped tar holding the report file exactly as the
    audit script wrote it, so decoding it removes every terminal artefact. Any
    whitespace the terminal or the chat client inserted is stripped before
    decoding. If no pack is present the raw text is used and unwrap() does what
    it can.
    """
    raw = open(path, encoding="utf-8", errors="replace").read()
    m = re.search(r"BEGIN REPORT PACK.*?===(.*?)===\s*END REPORT PACK", raw, re.S)
    if not m:
        return raw, False
    blob = re.sub(r"\s+", "", m.group(1))
    try:
        import base64 as b64, io as _io, tarfile
        tf = tarfile.open(fileobj=_io.BytesIO(b64.b64decode(blob)), mode="r:gz")
        for mem in tf.getmembers():
            if mem.name.startswith("report_") and mem.name.endswith(".txt"):
                return tf.extractfile(mem).read().decode("utf-8", "replace"), True
    except Exception as e:
        print("  (report pack in %s could not be decoded: %s; falling back to raw text)"
              % (os.path.basename(path), str(e)[:60]))
    return raw, False


def parse(path):
    host = os.path.splitext(os.path.basename(path))[0]
    text, packed = read_source(path)
    if packed:
        print("  %-14s decoded from REPORT PACK (no terminal damage)" % host)
    lines = unwrap(text)
    d = {"host": host, "date": "", "status": {}, "findings": [], "sec": {},
         "exposed": [], "accounts": [], "sshd": "", "uptime": "", "sessions": "",
         "fail_total": "", "btmp_days": "", "hist": [], "sudoers": "", "idle": []}
    cur = 0
    in_find = False
    for l in lines:
        m = re.match(r"^Server Security Review - (\S+)", l)
        if m:
            d["host"] = m.group(1)
        m = re.match(r"^Date:\s*(\d{4}-\d{2}-\d{2})", l)
        if m and not d["date"]:
            d["date"] = m.group(1)
        m = re.match(r"^Host:\s*(\S+)\s+Date:\s*(\S+)", l)
        if m:
            d["host"], d["date"] = m.group(1), m.group(2)
        m = re.match(r"^=== (\d)\.", l)
        if m:
            cur = int(m.group(1)); in_find = False
            d["sec"].setdefault(cur, [])
            continue
        if l.startswith("=== SUMMARY"):
            cur = 0; continue
        m = re.match(r"^(\d)\s+\S.*\s([OXΔ△])$", l)
        if m:
            d["status"][int(m.group(1))] = m.group(2).replace("Δ", "△")
            continue
        if l.startswith("Findings ("):
            in_find = True; continue
        if in_find:
            m = re.match(r"^\d+\.\s*(.+)$", l)
            if m:
                d["findings"].append(m.group(1).strip())
                continue
            if l.startswith(("Report:", "Summary:", "Evidence:", "Reminder:")):
                in_find = False
        if cur:
            d["sec"][cur].append(l)
        if l.startswith("ALL IFACES") or l.startswith("IFACE "):
            d["exposed"].append(l.split(None, 2)[-1] if l.startswith("ALL") else l)
        m = re.match(r"^\[info\] (\S+)\s+uid=(\d+)\s+shell=(\S+)\s+pwstatus=(\S+)\s+gecos=(.*)$", l)
        if m:
            d["accounts"].append(m.groups())
        if "sshd: port=" in l:
            d["sshd"] = l.split("[info]")[-1].strip()
        if l.startswith("[info] up ") or re.search(r"\bup \d+ days?", l):
            d["uptime"] = re.sub(r"^\[info\]\s*", "", l)
        m = re.search(r"Active sessions: (\d+)", l)
        if m:
            d["sessions"] = m.group(1)
        m = re.search(r"Failed logins in btmp: (\d+)\s*\(btmp begins: (.+?)\)", l)
        if m:
            d["fail_total"], d["btmp_begin"] = m.group(1), m.group(2)
        m = re.search(r"btmp retention: ~(\d+) days", l)
        if m:
            d["btmp_days"] = m.group(1)
        m = re.search(r"sudo group members: (.+)$", l)
        if m:
            d["sudoers"] = m.group(1)
        if re.match(r"^\d+:\S", l):
            d["hist"].append(l)
        if l.startswith("[info] IDLE"):
            d["idle"].append(re.sub(r"^\[info\] IDLE\s*", "", l))
    # Canonicalise before de-duplicating: two copies of the same finding that the
    # terminal broke at different columns are only equal once both are rebuilt
    # from their template.
    d["findings"] = list(dict.fromkeys(canonical(f) for f in d["findings"]))
    return d


def details(d, n):
    """The narrative Details cell for item n, assembled from the parsed facts."""
    if n == 1:
        s = "In use. %s" % (d["uptime"] or "uptime not captured")
        if d["sessions"]:
            s += "\n%s active session(s) at review time." % d["sessions"]
        return s
    if n == 2:
        s = "No Telnet or FTP listener found.\n" if any(
            "No Telnet/FTP" in x for x in d["sec"].get(2, [])) else ""
        if d["sshd"]:
            s += d["sshd"] + "\n"
        ex = [e for e in d["exposed"]]
        s += "%d listening socket(s) on all interfaces:\n" % len(ex)
        s += "\n".join("  " + e for e in ex[:24])
        if len(ex) > 24:
            s += "\n  ... %d more" % (len(ex) - 24)
        return s
    if n == 3:
        s = "%d accounts with UID >= 1000:\n" % len(d["accounts"])
        s += "\n".join("  %s(%s) shell=%s pw=%s gecos=%s" % (a[0], a[1], a[2], a[3], a[4])
                       for a in d["accounts"])
        if d["sudoers"]:
            s += "\nsudo group: " + d["sudoers"]
        return s
    if n == 4:
        s = ""
        for f in d["findings"]:
            if f.startswith("4: account(s)"):
                s += f[3:] + "\n"
        if d["idle"]:
            s += "Last logins:\n" + "\n".join("  " + x for x in d["idle"])
        exp = [f for f in d["findings"] if "no password expiry" in f]
        if exp:
            s += "\nNo password expiry configured on %d account(s)." % len(exp)
        return s.strip() or "No idle account beyond the threshold."
    if n == 5:
        s = "Failed logins in btmp: %s" % (d.get("fail_total") or "0")
        if d.get("btmp_begin"):
            s += " (btmp begins %s)" % d["btmp_begin"]
        if d["btmp_days"]:
            s += "\nbtmp retention: ~%s days." % d["btmp_days"]
        body = [x for x in d["sec"].get(5, []) if re.match(r"^\d+\s+\S+$", x)]
        if body:
            s += "\nBreakdown:\n" + "\n".join("  " + b for b in body[:12])
        return s
    if n == 6:
        s = ""
        for x in d["sec"].get(6, []):
            if x.startswith("[OK]") or "bash_history" in x:
                s += re.sub(r"^\[(OK|WARN)\]\s*", "", x) + "\n"
                break
        if d["hist"]:
            s += "Commands requiring explanation (%d):\n" % len(d["hist"])
            s += "\n".join("  " + h for h in d["hist"][:20])
        if any("HISTTIMEFORMAT" in f for f in d["findings"]):
            s += "\nHISTTIMEFORMAT is not set, so history has no timestamps."
        return s.strip()
    return ""


def classify(text):
    for rx, risk, sev, act in RULES:
        if re.search(rx, text, re.I):
            return risk, sev, act
    return "Needs review.", "Low", "Assess and record the outcome."


# ---- styling ---------------------------------------------------------------
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE = Font(bold=True, size=14)
SUB = Font(italic=True, size=9, color="555555")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
SEV_FILL = {"High": PatternFill("solid", fgColor="F8CBAD"),
            "Medium": PatternFill("solid", fgColor="FFE699"),
            "Low": PatternFill("solid", fgColor="E2EFDA")}
ST_FILL = {"O": PatternFill("solid", fgColor="E2EFDA"),
           "△": PatternFill("solid", fgColor="FFE699"),
           "X": PatternFill("solid", fgColor="F8CBAD")}


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.border, cell.alignment = H_FILL, H_FONT, BOX, CENTER


def build(hosts, out):
    wb = Workbook()
    today = dt.date.today()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Server Security Review - Summary (%d servers)" % len(hosts)
    ws["A1"].font = TITLE
    ws["A2"] = "Review date: %s    Generated from the audit script output by make_report.py" % today
    ws["A2"].font = SUB
    ws["A3"] = "Status legend:  O = compliant    /    △ = compliant with observations    /    X = non-compliant"
    ws["A3"].font = SUB
    header(ws, 5, ["No", "Category", "Check item"] + [h["host"] for h in hosts] + ["Common issue / note"])
    for n in range(1, 7):
        r = 5 + n
        ws.cell(row=r, column=1, value=n).alignment = CENTER
        ws.cell(row=r, column=2, value=ITEMS[n - 1][0]).alignment = WRAP
        ws.cell(row=r, column=3, value=ITEMS[n - 1][1]).alignment = WRAP
        for i, h in enumerate(hosts):
            st = h["status"].get(n, "?")
            c = ws.cell(row=r, column=4 + i, value=st)
            c.alignment = CENTER
            if st in ST_FILL:
                c.fill = ST_FILL[st]
        worst = [h["host"] for h in hosts if h["status"].get(n) == "X"]
        note = "Non-compliant on: %s" % ", ".join(worst) if worst else \
               ("Observations on %d host(s)." % sum(1 for h in hosts if h["status"].get(n) == "△")
                if any(h["status"].get(n) == "△" for h in hosts) else "Compliant on all hosts.")
        ws.cell(row=r, column=4 + len(hosts), value=note).alignment = WRAP
        for col in range(1, 5 + len(hosts)):
            ws.cell(row=r, column=col).border = BOX
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 46
    for i in range(len(hosts)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 12
    ws.column_dimensions[get_column_letter(4 + len(hosts))].width = 52

    # ---- per-host sheets ----
    for h in hosts:
        s = wb.create_sheet(h["host"][:31])
        s["A1"] = "Server Security Review - %s" % h["host"]
        s["A1"].font = TITLE
        s["A2"] = "Reviewed %s. Findings: %d." % (h["date"] or today, len(h["findings"]))
        s["A2"].font = SUB
        s["A3"] = "Status legend:  O = compliant  /  △ = compliant with observations  /  X = non-compliant"
        s["A3"].font = SUB
        header(s, 5, ["No", "Category", "Check item", "Details", "Status", "Evidence"])
        for n in range(1, 7):
            r = 5 + n
            s.cell(row=r, column=1, value=n).alignment = CENTER
            s.cell(row=r, column=2, value=ITEMS[n - 1][0]).alignment = WRAP
            s.cell(row=r, column=3, value=ITEMS[n - 1][1]).alignment = WRAP
            s.cell(row=r, column=4, value=details(h, n)).alignment = WRAP
            st = h["status"].get(n, "?")
            c = s.cell(row=r, column=5, value=st)
            c.alignment = CENTER
            if st in ST_FILL:
                c.fill = ST_FILL[st]
            s.cell(row=r, column=6, value="").alignment = WRAP
            for col in range(1, 7):
                s.cell(row=r, column=col).border = BOX
            s.row_dimensions[r].height = 120
        s.cell(row=13, column=1,
               value="Evidence column left blank for terminal screenshots; raw output is in evidence/ on the host.").font = SUB
        for col, w in zip("ABCDEF", (5, 26, 34, 88, 8, 26)):
            s.column_dimensions[col].width = w

    # ---- findings ----
    fs = wb.create_sheet("Findings & Actions")
    header(fs, 1, ["ID", "Server", "Ref No", "Finding", "Risk", "Severity",
                   "Recommended action", "Owner", "Target date"])
    rows = []
    for h in hosts:
        for f in h["findings"]:
            m = re.match(r"^(\d):\s*(.+)$", f)
            ref, text = (m.group(1), m.group(2)) if m else ("", f)
            risk, sev, act = classify(text)
            rows.append([h["host"], ref, text, risk, sev, act])
    rows.sort(key=lambda r: (SEV_ORDER[r[4]], r[0], r[1]))
    for i, r in enumerate(rows, 1):
        fs.cell(row=i + 1, column=1, value="F-%02d" % i).alignment = CENTER
        for j, v in enumerate(r, 2):
            fs.cell(row=i + 1, column=j, value=v).alignment = WRAP
        fs.cell(row=i + 1, column=6).fill = SEV_FILL[r[4]]
        fs.cell(row=i + 1, column=6).alignment = CENTER
        fs.cell(row=i + 1, column=8, value="")
        fs.cell(row=i + 1, column=9,
                value=str(today + dt.timedelta(days=SEV_DAYS[r[4]]))).alignment = CENTER
        for col in range(1, 10):
            fs.cell(row=i + 1, column=col).border = BOX
    for col, w in zip("ABCDEFGHI", (7, 14, 8, 62, 52, 10, 62, 14, 13)):
        fs.column_dimensions[col].width = w
    fs.freeze_panes = "A2"

    # ---- evidence commands ----
    es = wb.create_sheet("Evidence Commands")
    header(es, 1, ["Ref No", "Applies to", "Purpose", "Command", "Collected"])
    EV = [
        ("1", "All", "Uptime, load and session count", "uptime; who"),
        ("2", "All", "Listening ports with process names", "sudo ss -nltup"),
        ("2", "All", "Effective sshd configuration", "sudo sshd -T | grep -Ei 'permitrootlogin|passwordauthentication|^port'"),
        ("2", "NFS hosts", "NFS export scope", "cat /etc/exports; sudo exportfs -v; showmount -e localhost"),
        ("3", "All", "Accounts with UID >= 1000", "awk -F: '$3>=1000 && $3<65534 {print $1, $3, $5}' /etc/passwd"),
        ("3", "All", "Privileged group membership", "getent group sudo adm wheel"),
        ("3", "All", "Installed SSH keys", "sudo find /root /home -name authorized_keys -exec ls -la {} \\;"),
        ("4", "All", "Last login per account", "lastlog"),
        ("4", "All", "Password ageing per account", "sudo chage -l <user>"),
        ("5", "All", "Failed login attempts", "sudo lastb"),
        ("5", "All", "Auth log review", "sudo grep -E 'Failed password|Invalid user|Accepted ' /var/log/auth.log"),
        ("5", "All", "Log retention", "sudo ls -la /var/log/btmp /var/log/wtmp /var/log/auth.log*"),
        ("6", "All", "Root shell history", "sudo cat /root/.bash_history"),
        ("6", "All", "Scheduled tasks", "sudo crontab -l; sudo systemctl list-timers --all"),
        ("6", "All", "History timestamp setting", "sudo grep -r HISTTIMEFORMAT /etc/profile /etc/profile.d/"),
    ]
    for i, (ref, applies, purpose, cmd) in enumerate(EV, 2):
        for j, v in enumerate((ref, applies, purpose, cmd, "Yes"), 1):
            c = es.cell(row=i, column=j, value=v)
            c.alignment = WRAP if j in (3, 4) else CENTER
            c.border = BOX
    for col, w in zip("ABCDE", (8, 14, 40, 78, 11)):
        es.column_dimensions[col].width = w
    es.freeze_panes = "A2"

    wb.save(out)
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("indir", nargs="?", default=os.path.join(os.path.dirname(__file__), "outputs"))
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    files = sorted(glob.glob(os.path.join(a.indir, "*.txt")))
    if not files:
        raise SystemExit("no *.txt in %s" % a.indir)
    hosts = [parse(f) for f in files]
    out = a.out or os.path.join(os.path.dirname(__file__),
                                "server_security_review_%dservers_%s.xlsx"
                                % (len(hosts), dt.datetime.now().strftime("%Y%m%d_%H%M")))
    n = build(hosts, out)
    for h in hosts:
        print("  %-14s status=%s  findings=%d  accounts=%d  exposed=%d"
              % (h["host"], "".join(h["status"].get(i, "?") for i in range(1, 7)),
                 len(h["findings"]), len(h["accounts"]), len(h["exposed"])))
    print("\n%d hosts, %d findings -> %s" % (len(hosts), n, out))


if __name__ == "__main__":
    main()
